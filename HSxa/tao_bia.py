# -*- coding: utf-8 -*-
"""
Tạo bìa hồ sơ lưu trữ (.docx) hàng loạt từ dữ liệu Excel (.xlsx).

Cách hoạt động
--------------
- Dùng chính tệp BÌA mẫu (.docx) làm khuôn: giữ NGUYÊN khổ A4, lề, text box,
  viền khung, font, cỡ chữ và các tab stop của mẫu.
- Mỗi dòng dữ liệu trong Excel -> 1 bìa, mỗi bìa nằm trên 1 trang A4 riêng.
- Các đoạn/ô giá trị trong mẫu được định vị theo NHÃN cố định của bìa (vd
  "HỒ SƠ", "Từ ngày:", "Gồm:", "Phông số", dòng gạch ngang phân cách...)
  chứ KHÔNG theo một chuỗi ví dụ cụ thể nào. Nhờ vậy script hoạt động đúng
  cả khi mẫu đã điền sẵn ví dụ (vd Sample.docx) lẫn khi mẫu để TRẮNG hoàn
  toàn các ô giá trị (mẫu dùng in hàng loạt thực tế, vd "Bìa HS.docx").

Cách dùng
---------
    python tao_bia.py
hoặc tùy biến đường dẫn:
    python tao_bia.py --mau "Bìa HS.docx" --data "Dữ liệu bìa.xlsx" --out "Bia_ho_so.docx"

Cấu trúc cột Excel (Sheet đầu tiên, có thể có dòng tiêu đề):
    A: Phông số
    B: Mục lục số
    C: Hồ sơ số
    D: Trích yếu nội dung
    E: Ngày tháng bắt đầu, kết thúc   (vd 16/02/2006-19/02/2008)
    F: Số tờ
    G: Thời hạn bảo quản              (vd Vĩnh viễn)
    H: Tên phông                     (dòng đầu bìa, vd "ỦY BAN NHÂN DÂN XÃ ...")

Phông số / Mục lục số (cột A/B) đọc theo TỪNG DÒNG; nếu 1 dòng bỏ trống thì
dùng giá trị mặc định nhập từ giao diện/CLI (--phong-so/--muc-luc) nếu có.
"""

import argparse
import copy
import os
import re
import sys

from lxml import etree
import openpyxl

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
WP14 = "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing"
XML_SPACE = "http://www.w3.org/XML/1998/namespace"


def w(tag):
    return "{%s}%s" % (W, tag)


def qn(ns, tag):
    return "{%s}%s" % (ns, tag)


def runs_text(p):
    """Văn bản hiển thị của 1 đoạn (gộp các <w:t>)."""
    return "".join(t.text or "" for t in p.iter(w("t")))


def set_run_text(run, text):
    """Đặt lại văn bản cho 1 run, giữ nguyên định dạng (rPr). Xóa hết
    <w:t>/<w:tab>/<w:br> cũ rồi chèn văn bản mới; nếu text chứa ký tự xuống
    dòng (\\n) thì chèn <w:br/> để Word ngắt dòng đúng (vd dòng phông ở cột
    H nhập nhiều dòng bằng Alt+Enter trong Excel)."""
    for child in list(run):
        if child.tag in (w("t"), w("tab"), w("br")):
            run.remove(child)
    lines = str(text).replace("\r\n", "\n").replace("\r", "\n").split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            etree.SubElement(run, w("br"))
        t = etree.SubElement(run, w("t"))
        t.set(qn(XML_SPACE, "space"), "preserve")
        t.text = line


def set_run_text_keep_tab(run, text):
    """Như set_run_text, nhưng GIỮ NGUYÊN <w:tab/> đã có trong run (chỉ xóa
    <w:t>/<w:br>). Dùng cho các ô giá trị mà mẫu gộp chung dấu TAB canh cột
    với chỗ trống nhập giá trị vào cùng 1 run (thường gặp khi mẫu để TRẮNG
    ô giá trị: Word dồn TAB và chỗ trống văn bản vào chung 1 run)."""
    for child in list(run):
        if child.tag in (w("t"), w("br")):
            run.remove(child)
    tab = run.find(w("tab"))
    lines = str(text).replace("\r\n", "\n").replace("\r", "\n").split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            br = etree.Element(w("br"))
            (tab.addprevious(br) if tab is not None else run.append(br))
        t = etree.Element(w("t"))
        t.set(qn(XML_SPACE, "space"), "preserve")
        t.text = line
        (tab.addprevious(t) if tab is not None else run.append(t))


def clear_paragraph_runs(p):
    """Xóa hết <w:r> hiện có của đoạn p (giữ nguyên <w:pPr>)."""
    for r in p.findall(w("r")):
        p.remove(r)


def add_run(p, rpr_template, text):
    """Thêm 1 run mới vào cuối đoạn p, sao chép định dạng (rPr) từ
    rpr_template (có thể None) rồi đặt văn bản."""
    r = etree.SubElement(p, w("r"))
    if rpr_template is not None:
        r.append(copy.deepcopy(rpr_template))
    set_run_text(r, text)
    return r


def first_run_rpr(p):
    for r in p.findall(w("r")):
        rpr = r.find(w("rPr"))
        if rpr is not None:
            return rpr
    return None


def first_bold_run_rpr(p):
    for r in p.findall(w("r")):
        rpr = r.find(w("rPr"))
        if rpr is not None and rpr.find(w("b")) is not None:
            return rpr
    return None


def strip_bold_rpr(rpr):
    """Sao chép rpr, bỏ định dạng ĐẬM - dùng khi cả đoạn mẫu vốn toàn chữ
    đậm nhưng cần in giá trị bằng chữ THƯỜNG (vd ngày tháng, số tờ)."""
    if rpr is None:
        return None
    new_rpr = copy.deepcopy(rpr)
    for tag in (w("b"), w("bCs")):
        el = new_rpr.find(tag)
        if el is not None:
            new_rpr.remove(el)
    return new_rpr


def set_paragraph_text(p, text, fallback_rpr):
    """Đặt lại TOÀN BỘ văn bản của đoạn p bằng 1 run duy nhất. Giữ định
    dạng (rPr) của run đầu tiên nếu đoạn đã có sẵn chữ (mẫu điền ví dụ như
    Sample.docx); nếu đoạn hoàn toàn trống (mẫu để trắng, vd Bìa HS.docx)
    thì dùng fallback_rpr."""
    rpr = first_run_rpr(p)
    clear_paragraph_runs(p)
    add_run(p, rpr if rpr is not None else fallback_rpr, text)


def insert_blank_paragraph_before(p, ppr_source=None):
    """Chèn 1 đoạn TRỐNG (không có run) ngay TRƯỚC đoạn p, sao chép định
    dạng đoạn (pPr) từ ppr_source (mặc định lấy của chính p) - dùng để đẩy
    khối nội dung phía dưới xuống thấp hơn một chút trong khung.

    Lưu ý: pPr của bản thân p có thể mang cỡ chữ RẤT lớn (vd đoạn "HỒ SƠ"
    dùng cỡ 75pt cho nhãn tiêu đề) - nếu sao chép nguyên, dòng trống mới sẽ
    cao bất thường. Nên truyền ppr_source là 1 đoạn trống có sẵn gần đó
    (cỡ chữ vừa phải) để chỉ đẩy xuống "một chút" thay vì nhảy cỡ lớn."""
    new_p = etree.Element(w("p"))
    ppr = (ppr_source if ppr_source is not None else p).find(w("pPr"))
    if ppr is not None:
        new_p.append(copy.deepcopy(ppr))
    p.addprevious(new_p)


def find_paragraph(txbx, pred):
    """Đoạn ĐẦU TIÊN trong text box có văn bản thỏa `pred(text)`."""
    for p in txbx.findall(w("p")):
        if pred(runs_text(p)):
            return p
    return None


def _mk_rpr(sz):
    """Tạo <w:rPr> đậm, cỡ chữ sz (nửa-điểm) - dùng làm định dạng dự phòng
    khi đoạn mẫu hoàn toàn trống (không còn run/định dạng nào để kế thừa).
    Giá trị lấy từ chính bìa ví dụ (Sample.docx)."""
    rpr = etree.Element(w("rPr"))
    etree.SubElement(rpr, w("b"))
    etree.SubElement(rpr, w("bCs"))
    sz_el = etree.SubElement(rpr, w("sz"))
    sz_el.set(w("val"), str(sz))
    szcs_el = etree.SubElement(rpr, w("szCs"))
    szcs_el.set(w("val"), str(sz))
    return rpr


# Định dạng dự phòng cho 2 dòng đầu bìa (nhãn "PHÔNG LƯU TRỮ" + tên phông)
# và dòng trích yếu nội dung, dùng khi mẫu để các đoạn này hoàn toàn trống.
FALLBACK_RPR_ORG = _mk_rpr(30)
FALLBACK_RPR_TITLE = _mk_rpr(36)

DASH_RE = re.compile(r"^-{5,}$")
ORG_LABEL = "PHÔNG LƯU TRỮ"


def strip_org_label(org_val):
    """Bỏ tiền tố 'PHÔNG LƯU TRỮ' khỏi tên phông (cột H) nếu người nhập đã
    gõ sẵn trong Excel (tên phông lưu trữ chính thức thường có dạng đầy đủ
    'Phông lưu trữ ...') — tránh lặp lại với nhãn "PHÔNG LƯU TRỮ" mà bìa đã
    tự in ở dòng trên."""
    s = (org_val or "").strip()
    if s.upper().startswith(ORG_LABEL):
        rest = s[len(ORG_LABEL):].strip(" :-–—")
        return rest if rest else s
    return s


def set_field_before_tab(p, value):
    """Gộp `value` vào run ĐẬM đầu tiên nằm TRƯỚC (hoặc gộp chung với) run
    có TAB đầu tiên của đoạn p, xóa văn bản các run đậm còn lại trong nhóm.

    Dùng cho 'Phông số', 'Mục lục số', 'Hồ sơ số' — giá trị trong mẫu có
    thể là 1 run đậm độc lập ('01'), tách thành nhiều run đậm ('0'+'1'),
    hoặc (khi mẫu để TRẮNG) gộp luôn TAB canh cột vào chung run với chỗ
    trống giá trị."""
    bold_runs = []
    for r in p.findall(w("r")):
        rpr = r.find(w("rPr"))
        is_bold = rpr is not None and rpr.find(w("b")) is not None
        txt = "".join(t.text or "" for t in r.iter(w("t")))
        has_tab = r.find(w("tab")) is not None
        if is_bold and (txt.strip() or has_tab):
            bold_runs.append(r)
        if has_tab:
            break
    if not bold_runs:
        return False
    for i, r in enumerate(bold_runs):
        val = value if i == 0 else ""
        if r.find(w("tab")) is not None:
            set_run_text_keep_tab(r, val)
        else:
            set_run_text(r, val)
    return True


def set_field_after_tab(p, value):
    """Đặt `value` vào run ĐẬM đầu tiên SAU dấu TAB đầu tiên của đoạn p —
    dùng cho 'Thời hạn bảo quản' nằm cùng dòng với 'Mục lục số', sau dấu
    TAB canh cột. Ưu tiên run đã có sẵn chữ (mẫu điền ví dụ); nếu tất cả
    trống (mẫu trắng) thì lấy run đậm CUỐI CÙNG của đoạn (chỗ trống gộp
    chung TAB cuối, chờ nhập giá trị)."""
    seen_tab = False
    candidates = []
    for r in p.findall(w("r")):
        if seen_tab:
            rpr = r.find(w("rPr"))
            if rpr is not None and rpr.find(w("b")) is not None:
                candidates.append(r)
        if r.find(w("tab")) is not None:
            seen_tab = True
    if not candidates:
        return False
    target = next((r for r in candidates
                   if "".join(t.text or "" for t in r.iter(w("t"))).strip()),
                  candidates[-1])
    if target.find(w("tab")) is not None:
        set_run_text_keep_tab(target, value)
    else:
        set_run_text(target, value)
    return True


def fill_textbox(txbx, row, phong_so=None, muc_luc=None):
    """Điền dữ liệu cho 1 text box (1 bìa).

    Định vị theo NHÃN/CẤU TRÚC đoạn (không theo 1 chuỗi ví dụ cụ thể) nên
    hoạt động đúng cả với mẫu đã điền ví dụ (Sample.docx) lẫn mẫu để TRẮNG
    hoàn toàn các ô giá trị (mẫu in thật, vd "Bìa HS.docx")."""
    paras = txbx.findall(w("p"))
    texts = [runs_text(p) for p in paras]

    # --- Tên phông (2 dòng đầu bìa: nhãn "PHÔNG LƯU TRỮ" + tên phông) ---
    # Định vị theo VỊ TRÍ tương đối so với dòng gạch ngang phân cách ngay
    # dưới -> bền vững kể cả khi 2 đoạn này hoàn toàn trống trong mẫu.
    for i, t in enumerate(texts):
        if i >= 2 and DASH_RE.match(t.strip()):
            set_paragraph_text(paras[i - 2], ORG_LABEL, FALLBACK_RPR_ORG)
            set_paragraph_text(paras[i - 1], strip_org_label(row["org"]),
                               FALLBACK_RPR_ORG)
            break

    # --- Trích yếu nội dung: đoạn ngay sau nhãn "HỒ SƠ" ---
    # Chèn thêm 1 dòng trống ngay trước nhãn "HỒ SƠ" để cả khối (HỒ SƠ +
    # trích yếu + các dòng bên dưới) dịch xuống một chút, tránh dồn sát lên
    # đầu khung như trước. Lấy cỡ chữ từ đoạn trống sẵn có ngay phía trên
    # (không phải từ chính đoạn "HỒ SƠ", vốn cỡ chữ rất lớn).
    for i, t in enumerate(texts):
        if t.strip() == "HỒ SƠ" and i + 1 < len(paras):
            insert_blank_paragraph_before(
                paras[i], paras[i - 1] if i > 0 else None)
            set_paragraph_text(paras[i + 1], row["title"], FALLBACK_RPR_TITLE)
            break

    # --- Ngày bắt đầu / kết thúc: đoạn "Từ ngày: ... Đến ngày: ..." ---
    # Nhãn giữ định dạng gốc của mẫu (đậm); riêng giá trị ngày tháng in
    # THƯỜNG dù mẫu để cả đoạn đậm.
    p_date = find_paragraph(
        txbx, lambda t: "Từ" in t and "ngày" in t and "Đến" in t)
    if p_date is not None:
        rpr_label = first_run_rpr(p_date)
        rpr_value = strip_bold_rpr(rpr_label)
        clear_paragraph_runs(p_date)
        add_run(p_date, rpr_label, "Từ ngày: ")
        add_run(p_date, rpr_value, row["start"])
        add_run(p_date, rpr_label, "    Đến ngày: ")
        add_run(p_date, rpr_value, row["end"])

    # --- Số tờ: đoạn "Gồm:  NN tờ " ---
    # Số tờ in đậm; chữ "tờ" in THƯỜNG.
    p_gom = find_paragraph(txbx, lambda t: t.strip().startswith("Gồm"))
    if p_gom is not None:
        rpr_label = first_run_rpr(p_gom)
        rpr_bold = first_bold_run_rpr(p_gom)
        if rpr_bold is None:
            rpr_bold = rpr_label
        clear_paragraph_runs(p_gom)
        add_run(p_gom, rpr_label, "Gồm:  ")
        add_run(p_gom, rpr_bold, "%s" % row["so_to"])
        add_run(p_gom, rpr_label, " tờ ")

    # --- Phông số / Mục lục số / Thời hạn bảo quản / Hồ sơ số ---
    # (nhãn + dấu TAB canh cột - giữ cơ chế set_field_before/after_tab)
    ps = row.get("phong_so") or phong_so
    ml = row.get("muc_luc") or muc_luc
    for p in paras:
        t = runs_text(p)
        if "Phông" in t:
            if ps:
                set_field_before_tab(p, ps)
        elif "ục lục số" in t:
            if ml:
                set_field_before_tab(p, ml)
            set_field_after_tab(p, row["thbq"])
        elif "Hồ sơ số" in t:
            set_field_before_tab(p, row["ho_so_so"])


def normalize_date(token):
    """Chuẩn hóa 1 mốc ngày:
       - Đủ ngày/tháng/năm -> dd/mm/yyyy (ngày & tháng 2 chữ số, năm 4 chữ số)
       - Chỉ tháng/năm (vd '8/2019') -> mm/yyyy (tháng 2 chữ số, năm 4 chữ số)
       - Chỉ năm           -> yyyy
       - Dạng dính liền 'ddmmyyyy' (vd '22102020') -> dd/mm/yyyy
    """
    s = (token or "").strip().replace(" ", "")
    if not s:
        return ""

    def y4(y):
        return y if len(y) >= 4 else y.zfill(4)

    # 8 chữ số dính liền: ddmmyyyy
    if re.fullmatch(r"\d{8}", s):
        return "%s/%s/%s" % (s[0:2], s[2:4], s[4:8])

    parts = [p for p in s.split("/") if p != ""]
    if len(parts) == 3:                       # ngày / tháng / năm
        d, m, y = parts
        return "%s/%s/%s" % (d.zfill(2), m.zfill(2), y4(y))
    if len(parts) == 2:                       # tháng / năm (không có ngày)
        m, y = parts
        return "%s/%s" % (m.zfill(2), y4(y))
    if len(parts) == 1:                       # chỉ năm
        return parts[0]
    return s


def split_dates(raw):
    """'16/02/2006-19/02/2008' -> ('16/02/2006', '19/02/2008').
    Chuẩn hóa định dạng; chịu được khoảng trắng và trường hợp chỉ có 1 ngày."""
    s = (raw or "").strip()
    if not s:
        return "", ""
    # Tách trên dấu '-' nằm GIỮA hai ngày (ngày dạng dd/mm/yyyy không chứa '-')
    parts = re.split(r"\s*-\s*", s)
    if len(parts) >= 2:
        return normalize_date(parts[0]), normalize_date(parts[-1])
    return normalize_date(s), ""


def cell(ws, r, c):
    v = ws.cell(r, c).value
    return "" if v is None else str(v).strip()


def is_header(ho_so_so, date_range, so_to):
    """Dòng tiêu đề: không có chữ số ở Hồ sơ số / Ngày tháng / Số tờ."""
    return not re.search(r"\d", ho_so_so + " " + date_range + " " + so_to)


def read_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    return read_rows_ws(wb.worksheets[0])


def read_rows_ws(ws):
    rows = []
    for r in range(1, ws.max_row + 1):
        a = cell(ws, r, 1)   # Phông số
        b = cell(ws, r, 2)   # Mục lục số
        c = cell(ws, r, 3)   # Hồ sơ số
        d = cell(ws, r, 4)   # Trích yếu nội dung
        e = cell(ws, r, 5)   # Ngày tháng bắt đầu, kết thúc
        f = cell(ws, r, 6)   # Số tờ
        g = cell(ws, r, 7)   # Thời hạn bảo quản
        h = cell(ws, r, 8)   # Tên phông (đầu bìa)
        # Bỏ qua dòng trống hoàn toàn và dòng tiêu đề
        if not any([a, b, c, d, e, f, g, h]):
            continue
        if is_header(c, e, f):
            continue
        start, end = split_dates(e)
        rows.append({
            "phong_so": a,
            "muc_luc": b,
            "ho_so_so": c,
            "title": d,
            "start": start,
            "end": end,
            "so_to": f,
            "thbq": g,
            "org": h,
        })
    return rows


def _bottom_block_index(txbx):
    """Chỉ số đoạn bắt đầu khối dưới (đoạn chứa 'Phông')."""
    for i, p in enumerate(txbx.findall(w("p"))):
        if "Phông" in "".join(t.text or "" for t in p.iter(w("t"))):
            return i
    return None


def _set_nofill(sppr):
    """Bỏ nền và viền của shape (đặt noFill)."""
    for tag in ("solidFill", "gradFill", "blipFill", "pattFill", "noFill"):
        e = sppr.find(qn(A, tag))
        if e is not None:
            sppr.remove(e)
    # Chèn noFill ngay sau prstGeom (đúng thứ tự schema)
    geom = sppr.find(qn(A, "prstGeom"))
    nofill = etree.Element(qn(A, "noFill"))
    if geom is not None:
        geom.addnext(nofill)
    else:
        sppr.append(nofill)
    ln = sppr.find(qn(A, "ln"))
    if ln is not None:
        for c in list(ln):
            ln.remove(c)
        etree.SubElement(ln, qn(A, "noFill"))


def restructure_cover(cover_p):
    """Tách bìa thành 2 text box phủ toàn trang chồng lên nhau:
       - Box 1: giữ khung viền + nền + nội dung trên (đến trước 'Phông số'),
         canh đỉnh (anchor='t').
       - Box 2: trong suốt, không viền, chỉ chứa khối dưới ('Phông số'...),
         canh ĐÁY (anchor='b') -> luôn cố định ở đáy trang.
    """
    run = cover_p.find(w("r"))
    ac = run.find(qn(MC, "AlternateContent"))
    if ac is None:
        return  # đã tái cấu trúc rồi hoặc mẫu khác

    # 1) Bỏ AlternateContent, chỉ giữ <w:drawing> trong mc:Choice (box 1)
    choice = ac.find(qn(MC, "Choice"))
    draw1 = copy.deepcopy(choice.find(w("drawing")))
    idx = list(run).index(ac)
    run.remove(ac)
    run.insert(idx, draw1)

    txbx1 = draw1.find(".//" + w("txbxContent"))
    bidx = _bottom_block_index(txbx1)
    if bidx is None:
        return

    # 2) Tạo box 2 từ bản sao box 1
    draw2 = copy.deepcopy(draw1)

    # Box 1: bỏ các đoạn từ 'Phông số' trở xuống
    for p in txbx1.findall(w("p"))[bidx:]:
        txbx1.remove(p)

    # Box 2: chỉ giữ các đoạn từ 'Phông số' trở xuống
    txbx2 = draw2.find(".//" + w("txbxContent"))
    for p in txbx2.findall(w("p"))[:bidx]:
        txbx2.remove(p)

    # Box 2: canh nội dung xuống đáy
    body2 = draw2.find(".//" + qn(WPS, "bodyPr"))
    if body2 is not None:
        body2.set("anchor", "b")
    # Box 2: bỏ nền + viền (để lộ khung của box 1)
    sp2 = draw2.find(".//" + qn(WPS, "spPr"))
    if sp2 is not None:
        _set_nofill(sp2)
    # Box 2: id/anchorId/z-order khác box 1 để hợp lệ và nằm trên
    anc2 = draw2.find(qn(WP, "anchor"))
    if anc2 is not None:
        anc2.set(qn(WP14, "anchorId"), "0CBB7FA1")
        anc2.set(qn(WP14, "editId"), "040BC5A1")
        rh = anc2.get("relativeHeight")
        if rh:
            anc2.set("relativeHeight", str(int(rh) + 4))
    doc2 = draw2.find(".//" + qn(WP, "docPr"))
    if doc2 is not None:
        try:
            doc2.set("id", str(int(doc2.get("id")) + 1))
        except (TypeError, ValueError):
            doc2.set("id", "777000001")
        doc2.set("name", "BiaDuoi")

    # 3) Thêm box 2 vào cùng đoạn (run mới) -> hai box phủ cùng vị trí
    run2 = etree.SubElement(cover_p, w("r"))
    run2.append(draw2)


def build_document_xml(doc_xml, rows, progress=None, phong_so=None, muc_luc=None):
    """Từ document.xml của mẫu + danh sách dòng -> document.xml mới (bytes).
    progress(done, total): hàm tùy chọn báo tiến độ sau mỗi bìa.
    phong_so / muc_luc: giá trị MẶC ĐỊNH dùng khi 1 dòng Excel bỏ trống cột
    Phông số (A) / Mục lục số (B)."""
    root = etree.fromstring(doc_xml)
    body = root.find(w("body"))

    # Tách các thành phần body: đoạn chứa text box (cover) và sectPr
    children = list(body)
    cover_p = None
    sectpr = None
    for ch in children:
        if ch.tag == w("p") and ch.find(".//" + w("drawing")) is not None:
            cover_p = ch
        elif ch.tag == w("sectPr"):
            sectpr = ch
    if cover_p is None or sectpr is None:
        raise ValueError("Không tìm thấy text box bìa hoặc sectPr trong tệp mẫu.")

    # Cố định khối dưới ở đáy trang (tách thành text box thứ 2 canh đáy)
    restructure_cover(cover_p)

    # Dọn sạch body, dựng lại: cover_1, [ngắt trang], cover_2, ..., sectPr
    for ch in children:
        body.remove(ch)

    for idx, row in enumerate(rows):
        cp = copy.deepcopy(cover_p)
        # Điền dữ liệu cho TẤT CẢ text box bên trong (Choice + Fallback)
        for txbx in cp.iter(w("txbxContent")):
            fill_textbox(txbx, row, phong_so=phong_so, muc_luc=muc_luc)
        # Từ bìa thứ 2: ngắt sang trang mới
        if idx > 0:
            ppr = cp.find(w("pPr"))
            if ppr is None:
                ppr = etree.Element(w("pPr"))
                cp.insert(0, ppr)
            etree.SubElement(ppr, w("pageBreakBefore"))
        body.append(cp)
        if progress is not None:
            progress(idx + 1, len(rows))

    body.append(sectpr)
    return etree.tostring(root, xml_declaration=True,
                          encoding="UTF-8", standalone=True)


def assemble_docx(mau_bytes, new_document_xml):
    """Sao chép toàn bộ docx mẫu (bytes), chỉ thay word/document.xml -> bytes."""
    import zipfile, io
    src = io.BytesIO(mau_bytes)
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = new_document_xml
            zout.writestr(item, data)
    return out.getvalue()


def generate_from_bytes(mau_bytes, data_bytes, progress=None,
                        phong_so=None, muc_luc=None):
    """Lõi dùng cho web: nhận bytes mẫu + bytes Excel -> (docx bytes, danh sách dòng).
    progress(done, total): hàm tùy chọn báo tiến độ realtime.
    phong_so / muc_luc: giá trị MẶC ĐỊNH dùng khi 1 dòng Excel bỏ trống cột
    Phông số (A) / Mục lục số (B)."""
    import io, zipfile
    rows = read_rows_ws(
        openpyxl.load_workbook(io.BytesIO(data_bytes), data_only=True).worksheets[0])
    if not rows:
        raise ValueError("Không đọc được dòng dữ liệu nào từ Excel.")
    doc_xml = zipfile.ZipFile(io.BytesIO(mau_bytes)).read("word/document.xml")
    new_doc = build_document_xml(doc_xml, rows, progress=progress,
                                 phong_so=phong_so, muc_luc=muc_luc)
    return assemble_docx(mau_bytes, new_doc), rows


def build(mau_path, data_path, out_path, phong_so=None, muc_luc=None):
    with open(mau_path, "rb") as f:
        mau_bytes = f.read()
    with open(data_path, "rb") as f:
        data_bytes = f.read()
    out_bytes, rows = generate_from_bytes(mau_bytes, data_bytes,
                                          phong_so=phong_so, muc_luc=muc_luc)
    with open(out_path, "wb") as f:
        f.write(out_bytes)
    print("Đã tạo %d bìa -> %s" % (len(rows), out_path))


def autofind(here, pattern, exclude=(), exclude_prefixes=()):
    import glob
    cands = [p for p in glob.glob(os.path.join(here, pattern))
             if os.path.basename(p) not in exclude
             and not os.path.basename(p).startswith(tuple(exclude_prefixes))]
    return cands[0] if cands else None


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description="Tạo bìa hồ sơ từ Excel.")
    ap.add_argument("--mau", default=None, help="Tệp .docx mẫu (mặc định: tự dò)")
    ap.add_argument("--data", default=None, help="Tệp .xlsx dữ liệu (mặc định: tự dò)")
    ap.add_argument("--out", default=os.path.join(here, "Bia_ho_so_ket_qua.docx"),
                    help="Tệp .docx kết quả")
    ap.add_argument("--phong-so", default=None,
                    help="Phông số mặc định (dùng khi cột A bỏ trống)")
    ap.add_argument("--muc-luc", default=None,
                    help="Mục lục số mặc định (dùng khi cột B bỏ trống)")
    args = ap.parse_args()

    # Tự dò tệp nếu không chỉ định (tránh nhầm tệp kết quả / tệp ví dụ /
    # tệp khóa tạm "~$..." Word tạo ra khi tệp đang được mở)
    mau = args.mau or autofind(here, "*.docx",
                               exclude={os.path.basename(args.out)},
                               exclude_prefixes=("Sample", "~$"))
    data = args.data or autofind(here, "*.xlsx", exclude_prefixes=("~$",))

    if not mau or not os.path.exists(mau):
        sys.exit("Không tìm thấy tệp .docx mẫu.")
    if not data or not os.path.exists(data):
        sys.exit("Không tìm thấy tệp .xlsx dữ liệu.")

    print("Mẫu :", os.path.basename(mau))
    print("Data:", os.path.basename(data))
    build(mau, data, args.out, phong_so=args.phong_so, muc_luc=args.muc_luc)


if __name__ == "__main__":
    main()
