"""
CHUYỂN ĐỔI EXCEL SANG WORD - MỤC LỤC XÃ
Phiên bản 5.0 - Không lưu trữ trên server, tự động tải tệp về máy

Tối ưu cho PythonAnywhere (dung lượng đĩa rất hạn chế):
  - KHÔNG ghi bất kỳ tệp nào lên đĩa: đọc Excel từ RAM, xuất Word ra RAM,
    trả thẳng về trình duyệt rồi giải phóng bộ nhớ.
  - Trình duyệt tự động tải tệp .docx xuống ngay khi xử lý xong.

Khác biệt so với app-ok.py:
  1. Bỏ toàn bộ hệ thống đăng nhập (Flask-Login) - vào thẳng trang chính
  2. Bỏ ô nhập "Năm hồ sơ" ở UI và dòng "Năm: ..." trong tệp Word đầu ra
  3. Đảo vị trí cột D và E:
       - Cột D: Tác giả tài liệu          (trước đây là Trích yếu nội dung)
       - Cột E: Tên loại và trích yếu nội dung văn bản  (trước đây là Tác giả)
"""

import io
import os
import re
import pandas as pd
from datetime import datetime
from flask import Flask, request, jsonify, send_file, render_template_string
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ==================== CẤU HÌNH ====================
# Không có thư mục làm việc - toàn bộ xử lý diễn ra trong bộ nhớ
A4_WIDTH_IN = 8.27
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB

# ==================== TIỆN ÍCH ====================
def log(msg):
    timestamp = datetime.now().strftime('%H:%M:%S')
    print(f"[{timestamp}] {msg}")

# ==================== HÀM XỬ LÝ WORD ====================
def inches_to_twips(inches: float) -> int:
    return int(round(inches * 1440))

def set_table_full_width_perfect(table, usable_width_inches: float):
    tbl = table._tbl

    if tbl.tblPr is not None:
        tbl.remove(tbl.tblPr)

    tblPr = OxmlElement('w:tblPr')
    tbl.insert(0, tblPr)

    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(inches_to_twips(usable_width_inches)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)

    tblInd = OxmlElement('w:tblInd')
    tblInd.set(qn('w:w'), '0')
    tblInd.set(qn('w:type'), 'dxa')
    tblPr.append(tblInd)

    tblCellMar = OxmlElement('w:tblCellMar')
    for side in ['top', 'left', 'bottom', 'right']:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), '0')
        el.set(qn('w:type'), 'dxa')
        tblCellMar.append(el)
    tblPr.append(tblCellMar)

    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'fixed')
    tblPr.append(tblLayout)

    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def set_cell_width_exact(cell, width_inches: float):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    for old in tcPr.findall(qn('w:tcW')):
        tcPr.remove(old)

    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(inches_to_twips(width_inches)))
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)

def extract_sort_key(key):
    if key is None:
        return (1_000_000_000, "")
    s = str(key).strip()
    m = re.search(r'(\d+)', s)
    if m:
        return (int(m.group(1)), s)
    return (1_000_000_000, s)

# ==================== HÀM XỬ LÝ CHÍNH ====================
def build_docx_in_memory(excel_bytes: bytes) -> io.BytesIO:
    """
    Đọc dữ liệu Excel từ bộ nhớ, dựng tài liệu Word và trả về BytesIO.
    Không đụng tới đĩa ở bất kỳ bước nào.
    """
    df = pd.read_excel(io.BytesIO(excel_bytes), header=None, dtype=str)

    if df.empty:
        raise ValueError("Tệp Excel không có dữ liệu")
    if df.shape[1] < 8:
        raise ValueError(
            f"Tệp Excel cần ít nhất 8 cột (A: Hồ sơ số, B-H: nội dung bảng), "
            f"hiện chỉ có {df.shape[1]} cột"
        )

    # Bỏ hàng tiêu đề
    df = df.iloc[1:].reset_index(drop=True)
    if df.empty:
        raise ValueError("Tệp Excel chỉ có hàng tiêu đề, không có dữ liệu")

    df[0] = df[0].ffill()

    raw_groups = list(df.groupby(df[0]))
    groups = sorted(raw_groups, key=lambda x: extract_sort_key(x[0]))
    if not groups:
        raise ValueError("Không tìm thấy hồ sơ nào trong tệp")

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(13)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    section = doc.sections[0]
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(0.4)
    usable_width_in = A4_WIDTH_IN - 0.8

    # Cột 4 (Tên loại và trích yếu nội dung văn bản) là cột rộng nhất,
    # cột 3 (Tác giả tài liệu) hẹp lại - đảo so với bản cũ
    fixed = {0: 0.37, 5: 0.37, 6: 0.55}
    flexible = {1: 1.5, 2: 1.7, 3: 2.1, 4: 5.8}

    total_fixed = sum(fixed.values())
    remaining = usable_width_in - total_fixed
    sum_weights = sum(flexible.values())
    flex_widths = {i: (flexible[i] / sum_weights) * remaining for i in flexible}

    final_widths = []
    for i in range(7):
        if i in fixed:
            final_widths.append(fixed[i])
        else:
            final_widths.append(flex_widths[i])

    # Dồn phần dư vào cột trích yếu (index 4)
    total = sum(final_widths)
    final_widths[4] += (usable_width_in - total)

    headers = [
        "Số TT", "Số, ký hiệu tài liệu", "Ngày tháng văn bản",
        "Tác giả tài liệu", "Tên loại và trích yếu nội dung văn bản", "Tờ số", "Ghi chú"
    ]

    total_groups = len(groups)
    for gi, (so_tep, group) in enumerate(groups):
        group = group.dropna(subset=[1]).reset_index(drop=True)

        p = doc.add_paragraph("MỤC LỤC VĂN BẢN", style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(14)

        p = doc.add_paragraph(f"Hồ sơ số: {so_tep}", style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(14)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        table = doc.add_table(rows=1, cols=7)
        table.style = 'Table Grid'
        table.allow_autofit = False
        table.autofit = False

        set_table_full_width_perfect(table, usable_width_in)

        for col_idx, width_in in enumerate(final_widths):
            for cell in table.columns[col_idx].cells:
                set_cell_width_exact(cell, width_in)

        hdr_cells = table.rows[0].cells
        for i, text in enumerate(headers):
            para = hdr_cells[i].paragraphs[0]
            # căn sát lề header
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1
            para.paragraph_format.keep_together = True
            # end- căn lề

            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            run = para.add_run(text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

        tr = table.rows[0]._tr
        trPr = tr.get_or_add_trPr()
        tblHeader = OxmlElement('w:tblHeader')
        trPr.append(tblHeader)

        # Dữ liệu lấy theo đúng thứ tự cột của tệp Excel đầu vào
        # (tệp đầu vào đã đổi cột D/E tương ứng nên giữ nguyên ánh xạ vị trí)
        for _, row in group.iterrows():
            cells = table.add_row().cells
            for i in range(7):
                val = str(row[i + 1]) if pd.notna(row[i + 1]) else ""
                para = cells[i].paragraphs[0]
                para.text = val

                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1
                para.paragraph_format.keep_together = True

                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                run = para.runs[0] if para.runs else para.add_run()
                run.font.name = 'Times New Roman'
                run.font.size = Pt(13)
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

                set_cell_width_exact(cells[i], final_widths[i])

        if gi < total_groups - 1:
            doc.add_page_break()

    # Xuất ra bộ nhớ thay vì lưu đĩa
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ==================== ROUTES ====================
@app.route("/")
def index():
    """Trang chính - vào thẳng, không cần đăng nhập"""
    return render_template_string(HTML)

@app.route("/convert", methods=["POST"])
def convert():
    """
    Nhận tệp Excel, chuyển đổi trong bộ nhớ và trả thẳng tệp Word về trình duyệt.
    Server không lưu lại bất cứ thứ gì.
    """
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"success": False, "error": "Chưa chọn tệp"}), 400
    if not file.filename.lower().endswith(".xlsx"):
        return jsonify({"success": False, "error": "Chỉ chấp nhận tệp .xlsx"}), 400

    # Tên tệp gốc chỉ dùng để đặt tên tải về, không dùng làm đường dẫn
    base_name = os.path.splitext(os.path.basename(file.filename))[0]
    out_name = f"{base_name}.docx"

    try:
        excel_bytes = file.read()
        buffer = build_docx_in_memory(excel_bytes)
    except ValueError as e:
        log(f"⚠️  {file.filename}: {e}")
        return jsonify({"success": False, "error": str(e)}), 400
    except Exception as e:
        log(f"❌ Lỗi khi xử lý {file.filename}: {e}")
        return jsonify({"success": False, "error": f"Lỗi khi xử lý tệp: {e}"}), 500

    size_kb = buffer.getbuffer().nbytes / 1024
    log(f"✅ {file.filename} → {out_name} ({size_kb:.0f} KB, không lưu đĩa)")

    return send_file(
        buffer,
        mimetype=DOCX_MIME,
        as_attachment=True,
        download_name=out_name,
        max_age=0,
    )

@app.route("/health")
def health():
    """Kiểm tra nhanh app còn sống"""
    return jsonify({"status": "ok", "storage": "none - xử lý trong bộ nhớ"})

# ==================== HTML TEMPLATE CHÍNH ====================
HTML = '''
<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Mục lục Xã - Chuyển đổi Excel sang Word</title>
<link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.min.css" rel="stylesheet">
<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/js/bootstrap.bundle.min.js"></script>
<style>
:root {
    --primary: #0d6efd;
    --success: #198754;
    --danger: #dc3545;
    --info: #0dcaf0;
}

body {
    background: linear-gradient(135deg, #f5f7fa 0%, #c3cfe2 100%);
    font-family: 'Segoe UI', Tahoma, sans-serif;
    min-height: 100vh;
    padding: 20px 0;
}

.container-card {
    background: white;
    border-radius: 20px;
    box-shadow: 0 15px 35px rgba(50, 50, 93, 0.1), 0 5px 15px rgba(0, 0, 0, 0.07);
    padding: 30px;
    margin: 0 auto;
    max-width: 1000px;
    position: relative;
    overflow: hidden;
}

.container-card::before {
    content: '';
    position: absolute;
    top: 0;
    left: 0;
    right: 0;
    height: 5px;
    background: linear-gradient(to right, var(--primary), var(--info));
}

h1 {
    color: var(--primary);
    font-weight: 700;
    margin-bottom: 10px;
    position: relative;
    padding-bottom: 15px;
}

h1::after {
    content: '';
    position: absolute;
    bottom: 0;
    left: 0;
    width: 60px;
    height: 3px;
    background: var(--primary);
}

.task-card {
    background: #f8f9fa;
    border-radius: 12px;
    border-left: 5px solid var(--info);
    padding: 15px;
    margin-bottom: 12px;
    transition: all 0.3s ease;
    animation: slideIn 0.3s ease-out;
}

.task-card.processing { border-left-color: var(--info); background: #e8f4fc; }
.task-card.completed { border-left-color: var(--success); background: #e8f5e9; }
.task-card.error { border-left-color: var(--danger); background: #fde8e8; }
.task-card.queued { border-left-color: #6c757d; background: #f8f9fa; }

@keyframes slideIn {
    from { opacity: 0; transform: translateY(-10px); }
    to { opacity: 1; transform: translateY(0); }
}

.progress-container {
    height: 30px;
    border-radius: 15px;
    overflow: hidden;
    background: #e9ecef;
    margin: 10px 0;
    box-shadow: inset 0 1px 3px rgba(0,0,0,0.2);
}

.progress-bar-custom {
    height: 100%;
    background: linear-gradient(90deg, var(--info), var(--primary));
    transition: width 0.5s ease;
    display: flex;
    align-items: center;
    justify-content: center;
    color: white;
    font-weight: 600;
    font-size: 0.9rem;
}

.log-box {
    background: #1a1a1a;
    color: #00ff88;
    font-family: 'Consolas', 'Monaco', monospace;
    padding: 15px;
    border-radius: 10px;
    height: 250px;
    overflow-y: auto;
    font-size: 0.9rem;
    border: 1px solid #333;
}

.log-box::-webkit-scrollbar { width: 8px; }
.log-box::-webkit-scrollbar-track { background: #2d2d2d; }
.log-box::-webkit-scrollbar-thumb { background: var(--primary); border-radius: 4px; }

.file-item {
    display: flex;
    justify-content: space-between;
    align-items: center;
    padding: 10px 15px;
    background: white;
    border: 1px solid #dee2e6;
    border-radius: 8px;
    margin-bottom: 8px;
    transition: all 0.2s;
}

.file-item:hover {
    transform: translateY(-2px);
    box-shadow: 0 4px 12px rgba(0,0,0,0.1);
    border-color: var(--primary);
}

#overlay {
    position: fixed;
    top: 0;
    left: 0;
    right: 0;
    bottom: 0;
    background: rgba(255, 255, 255, 0.92);
    backdrop-filter: blur(8px);
    display: flex;
    flex-direction: column;
    align-items: center;
    justify-content: center;
    z-index: 9999;
    opacity: 0;
    pointer-events: none;
    transition: opacity 0.3s ease;
}

#overlay.show { opacity: 1; pointer-events: all; }

.spinner {
    width: 60px;
    height: 60px;
    border: 5px solid rgba(13, 110, 253, 0.1);
    border-top-color: var(--primary);
    border-radius: 50%;
    animation: spin 1s linear infinite;
    margin-bottom: 20px;
}

@keyframes spin { to { transform: rotate(360deg); } }
</style>
</head>
<body>

<div id="overlay">
    <div class="spinner"></div>
    <h4 class="mt-3 text-primary" id="overlayText">Đang xử lý...</h4>
    <div class="progress mt-3" style="width: 200px;">
        <div id="overlayProgress" class="progress-bar progress-bar-striped progress-bar-animated" style="width: 0%"></div>
    </div>
</div>

<div class="container">
    <div class="container-card">
        <div class="text-center mb-4">
            <h1>🔄 Mục lục Xã: Excel → Word</h1>
            <p class="text-muted">Phiên bản 5.0 - Tự động tải về, server không lưu trữ</p>
            <span class="badge bg-success">🔒 Không lưu tệp trên server</span>
        </div>

        <div class="card border-0 shadow-sm mb-4">
            <div class="card-body">
                <h5 class="card-title mb-3">📤 Chọn tệp Excel</h5>
                <div class="row g-3">
                    <div class="col-md-9">
                        <label class="form-label fw-semibold">📁 Tệp Excel (.xlsx)</label>
                        <input id="fileInput" class="form-control" type="file" accept=".xlsx" multiple>
                        <div class="form-text">
                            Có thể chọn nhiều tệp. Tệp Word sẽ tự động tải xuống ngay khi chuyển đổi xong.
                        </div>
                    </div>
                    <div class="col-md-3 d-flex align-items-end">
                        <button id="uploadBtn" class="btn btn-primary w-100" onclick="startUpload()">🚀 Bắt đầu</button>
                    </div>
                </div>
                <div class="alert alert-info mt-3 mb-0 py-2 small">
                    ℹ️ Nếu chọn nhiều tệp, trình duyệt có thể hỏi <strong>"Cho phép tải nhiều tệp?"</strong> —
                    hãy bấm <strong>Cho phép</strong>.
                </div>
            </div>
        </div>

        <div class="card border-0 shadow-sm mb-4">
            <div class="card-body">
                <h5 class="card-title mb-3">📊 Tiến độ xử lý</h5>
                <div id="progressWrap" style="display: none;">
                    <div class="d-flex justify-content-between mb-2">
                        <span id="statusText" class="fw-semibold">Đang chuẩn bị...</span>
                        <span id="progressPercent" class="fw-bold text-primary">0%</span>
                    </div>
                    <div class="progress-container">
                        <div id="progressBar" class="progress-bar-custom" style="width: 0%">0%</div>
                    </div>
                </div>
                <div id="tasksContainer" class="mt-3">
                    <div class="text-center text-muted py-3">Chưa có tệp nào đang xử lý</div>
                </div>
            </div>
        </div>

        <div class="card border-0 shadow-sm mb-4">
            <div class="card-body">
                <h5 class="card-title mb-3">📄 Tệp đã chuyển đổi trong phiên này</h5>
                <div id="fileList">
                    <div class="text-center text-muted py-4">Chưa có tệp nào</div>
                </div>
                <div class="form-text mt-2">
                    Danh sách chỉ nằm trong trình duyệt của bạn. Tải lại trang là mất — server không giữ bản sao nào.
                </div>
            </div>
        </div>

        <div class="card border-0 shadow-sm">
            <div class="card-body">
                <h5 class="card-title mb-3">📝 Nhật ký hệ thống</h5>
                <div class="log-box" id="logBox">[Hệ thống] Đã sẵn sàng...</div>
                <div class="mt-2 text-end">
                    <button class="btn btn-sm btn-outline-secondary" onclick="clearLog()">Xóa log</button>
                </div>
            </div>
        </div>

        <div class="mt-4 text-center text-muted small">
            <p>© 2024 Tuấn IT - Phiên bản 5.0 - Mục lục Xã</p>
        </div>
    </div>
</div>

<div class="modal fade" id="doneModal" tabindex="-1">
    <div class="modal-dialog modal-dialog-centered">
        <div class="modal-content">
            <div class="modal-header bg-primary text-white">
                <h5 class="modal-title">🎉 Hoàn tất!</h5>
                <button type="button" class="btn-close btn-close-white" data-bs-dismiss="modal"></button>
            </div>
            <div class="modal-body text-center">
                <p id="doneModalText">Đã chuyển đổi xong và tải tệp về máy.</p>
                <div class="mt-3">
                    <button type="button" class="btn btn-primary" data-bs-dismiss="modal">Đóng</button>
                </div>
            </div>
        </div>
    </div>
</div>

<div class="toast-container position-fixed top-0 end-0 p-3">
    <div id="toast" class="toast" role="alert">
        <div class="toast-header">
            <strong class="me-auto" id="toastTitle">Thông báo</strong>
            <button type="button" class="btn-close" data-bs-dismiss="toast"></button>
        </div>
        <div class="toast-body" id="toastBody"></div>
    </div>
</div>

<script>
// Tiền tố đường dẫn: chạy độc lập -> "", khi gộp mount tại /mlxa -> "/mlxa".
const BASE = location.pathname.replace(/\\/+$/, "");
// Mỗi phần tử: {name, status, message}
let tasks = [];
// Giữ các blob URL để người dùng tải lại được, thu hồi khi rời trang
let objectUrls = [];

function showToast(title, message, type = 'info') {
    const toastEl = document.getElementById('toast');
    document.getElementById('toastTitle').textContent = title;
    document.getElementById('toastBody').textContent = message;
    const header = toastEl.querySelector('.toast-header');
    header.className = 'toast-header';
    if (type === 'success') header.classList.add('bg-success', 'text-white');
    else if (type === 'error') header.classList.add('bg-danger', 'text-white');
    else if (type === 'warning') header.classList.add('bg-warning', 'text-dark');
    else header.classList.add('bg-info', 'text-white');
    new bootstrap.Toast(toastEl).show();
}

function log(message) {
    const box = document.getElementById('logBox');
    const timestamp = new Date().toLocaleTimeString();
    box.innerHTML += `[${timestamp}] ${message}<br>`;
    box.scrollTop = box.scrollHeight;
}

function clearLog() {
    document.getElementById('logBox').innerHTML = '[Hệ thống] Log đã làm sạch<br>';
}

function setUploadingUI(isUploading, text) {
    const overlay = document.getElementById('overlay');
    document.getElementById('uploadBtn').disabled = isUploading;
    document.getElementById('fileInput').disabled = isUploading;
    if (text) document.getElementById('overlayText').textContent = text;
    overlay.classList.toggle('show', isUploading);
}

function renderTasks() {
    const container = document.getElementById('tasksContainer');
    if (tasks.length === 0) {
        container.innerHTML = '<div class="text-center text-muted py-3">Chưa có tệp nào đang xử lý</div>';
        return;
    }
    const labels = {
        processing: ['info', 'Đang xử lý'],
        completed: ['success', 'Hoàn thành'],
        error: ['danger', 'Lỗi'],
        queued: ['secondary', 'Đang chờ']
    };
    container.innerHTML = tasks.map(t => {
        const [color, label] = labels[t.status] || labels.queued;
        return `
        <div class="task-card ${t.status}">
            <div class="d-flex justify-content-between align-items-center">
                <div><strong>${t.name}</strong></div>
                <span class="badge bg-${color}">${label}</span>
            </div>
            ${t.message ? `<div class="mt-2 small">${t.message}</div>` : ''}
        </div>`;
    }).join('');
}

function setOverallProgress(done, total) {
    const pct = total ? Math.round(done / total * 100) : 0;
    document.getElementById('progressPercent').textContent = `${pct}%`;
    const bar = document.getElementById('progressBar');
    bar.style.width = `${pct}%`;
    bar.textContent = `${pct}%`;
    document.getElementById('statusText').textContent = `Đã xong ${done}/${total} tệp`;
    document.getElementById('overlayProgress').style.width = `${pct}%`;
}

// Kích hoạt tải xuống tự động
function triggerDownload(url, filename) {
    const a = document.createElement('a');
    a.href = url;
    a.download = filename;
    a.style.display = 'none';
    document.body.appendChild(a);
    a.click();
    a.remove();
}

function addFileToList(filename, url) {
    const fileList = document.getElementById('fileList');
    if (fileList.querySelector('.text-center')) fileList.innerHTML = '';
    const div = document.createElement('div');
    div.className = 'file-item';
    div.innerHTML = `
        <div>
            <span>📄 ${filename}</span>
            <div class="small text-muted">${new Date().toLocaleString()}</div>
        </div>
        <a class="btn btn-sm btn-outline-primary" download="${filename}">⬇️ Tải lại</a>
    `;
    div.querySelector('a').href = url;
    fileList.prepend(div);
}

async function startUpload() {
    const files = Array.from(document.getElementById('fileInput').files);
    if (!files.length) {
        showToast('Thiếu tệp', 'Vui lòng chọn tệp Excel', 'warning');
        return;
    }

    tasks = files.map(f => ({ name: f.name, status: 'queued', message: '' }));
    renderTasks();
    document.getElementById('progressWrap').style.display = 'block';
    setOverallProgress(0, files.length);
    setUploadingUI(true, 'Đang xử lý...');

    let okCount = 0;

    for (let i = 0; i < files.length; i++) {
        const file = files[i];
        tasks[i].status = 'processing';
        tasks[i].message = 'Đang chuyển đổi trên server...';
        renderTasks();
        setUploadingUI(true, `Đang xử lý ${i + 1}/${files.length}: ${file.name}`);

        const formData = new FormData();
        formData.append('file', file);

        try {
            const response = await fetch(BASE + '/convert', { method: 'POST', body: formData });
            const contentType = response.headers.get('Content-Type') || '';

            // Server trả JSON nghĩa là có lỗi; trả blob nghĩa là thành công
            if (!response.ok || contentType.includes('application/json')) {
                let err = 'Lỗi không xác định';
                try { err = (await response.json()).error || err; } catch (e) {}
                tasks[i].status = 'error';
                tasks[i].message = err;
                log(`❌ ${file.name}: ${err}`);
                showToast('Lỗi', `${file.name}: ${err}`, 'error');
            } else {
                const blob = await response.blob();
                const outName = file.name.replace(/\\.xlsx$/i, '') + '.docx';
                const url = URL.createObjectURL(blob);
                objectUrls.push(url);

                triggerDownload(url, outName);
                addFileToList(outName, url);

                okCount++;
                tasks[i].status = 'completed';
                tasks[i].message = `Đã tải xuống: ${outName} (${Math.round(blob.size / 1024)} KB)`;
                log(`✅ ${outName} — đã tải về máy`);
            }
        } catch (error) {
            tasks[i].status = 'error';
            tasks[i].message = 'Không kết nối được tới server';
            log(`❌ ${file.name}: lỗi mạng`);
            showToast('Lỗi', `${file.name}: lỗi mạng`, 'error');
        }

        renderTasks();
        setOverallProgress(i + 1, files.length);
    }

    setUploadingUI(false);

    if (okCount > 0) {
        document.getElementById('doneModalText').textContent =
            `Đã chuyển đổi ${okCount}/${files.length} tệp và tải về máy. Server không lưu bản sao nào.`;
        new bootstrap.Modal(document.getElementById('doneModal')).show();
        showToast('Thành công', `Đã tải xuống ${okCount} tệp`, 'success');
    }
}

// Giải phóng bộ nhớ blob khi rời trang
window.addEventListener('beforeunload', function () {
    objectUrls.forEach(u => URL.revokeObjectURL(u));
});

document.addEventListener('DOMContentLoaded', function () {
    log('Hệ thống sẵn sàng - chế độ không lưu trữ');
});
</script>
</body>
</html>
'''

# ==================== CHẠY SERVER ====================
# PythonAnywhere dùng WSGI nên chỉ cần trỏ tới biến `app` ở trên,
# khối dưới đây chỉ chạy khi khởi động thủ công trên máy cá nhân.
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    log("=" * 60)
    log("🚀 SERVER MỤC LỤC XÃ ĐÃ KHỞI ĐỘNG")
    log("💾 Chế độ: KHÔNG lưu trữ - xử lý trong bộ nhớ, tự động tải về")
    log(f"🔗 Truy cập: http://localhost:{port}/")
    log("=" * 60)

    app.run(host="0.0.0.0", port=port, debug=False, threaded=True)
